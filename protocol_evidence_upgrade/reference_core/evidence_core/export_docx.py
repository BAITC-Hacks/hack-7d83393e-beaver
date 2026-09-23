"""Render the saved snapshot without any new model call."""
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .audit import audit


def render(protocol: dict) -> bytes:
    report = audit(protocol)
    if report['status'] != 'PASS':
        raise ValueError('Экспорт отклонён: снимок не прошёл структурную проверку')
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.72)
    section.left_margin = section.right_margin = Inches(.75)
    normal = doc.styles['Normal']
    normal.font.name = 'DejaVu Sans'
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for name in ('Title','Heading 1','Heading 2'):
        doc.styles[name].font.name = 'DejaVu Sans'
        doc.styles[name].font.color.rgb = RGBColor.from_string('18363B')
    doc.core_properties.title = protocol['meeting']['title']
    doc.core_properties.author = 'Local protocol reference'
    doc.core_properties.subject = 'REFERENCE / SYNTHETIC DATA' if protocol['mode'] == 'FIXTURE' else 'Imported protocol snapshot'
    doc.add_heading('Протокол совещания', 0)
    doc.add_paragraph(protocol['meeting']['title'])
    p = doc.add_paragraph()
    p.add_run('СИНТЕТИЧЕСКИЙ ПРИМЕР. НЕ РЕЗУЛЬТАТ ASR/LLM.' if protocol['mode'] == 'FIXTURE' else 'ИМПОРТИРОВАННЫЙ СНИМОК. ПРОИСХОЖДЕНИЕ ASR НЕ ПРОВЕРЕНО.').bold = True
    # v1.0 has no protocol-level approval; do not infer approval of an entire document.
    doc.add_paragraph('Черновик. Структурная проверка не подтверждает смысловую правильность.')
    doc.add_paragraph(f"Начало: {protocol['meeting']['started_at']}\nЧасовой пояс: {protocol['meeting']['timezone']}")
    names = {p['id']: p['name'] for p in protocol['participants']}
    doc.add_paragraph('Участники: ' + ', '.join(names.values()))
    doc.add_heading('Краткое содержание', 1)
    for summary in protocol['summary']:
        doc.add_paragraph(summary['text'])
    if not protocol['summary']:
        doc.add_paragraph('Содержание не сформировано.')
    doc.add_heading('Поручения', 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Shading Accent 1'
    for cell, text in zip(table.rows[0].cells, ['Поручение', 'Исполнитель', 'Срок']):
        cell.text = text
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader')
    tr_pr.append(repeat)
    for task in protocol['tasks']:
        cells = table.add_row().cells
        cells[0].text = task['action'] + (' [отменено]' if task['dialogue_status'] == 'cancelled' else '')
        cells[1].text = names.get(task['assignee_id'], 'Не определён')
        cells[2].text = ((task['due_date'] or 'Не определён') + (' ' + task['due_time'] if task['due_time'] else ''))
    if not protocol['tasks']:
        doc.add_paragraph('Поручений нет.')
    if protocol['review_questions']:
        doc.add_heading('Вопросы перед утверждением', 1)
        for question in protocol['review_questions']:
            doc.add_paragraph(question['text'])
    doc.add_heading('Основания и история', 1)
    by_id = {u['id']: u for u in protocol['utterances']}
    for task in protocol['tasks']:
        doc.add_heading(task['action'], 2)
        for label, key in [('Действие','action'),('Исполнитель','assignee'),('Срок','deadline')]:
            refs = task['evidence'][key]
            doc.add_paragraph(label + ': ' + (', '.join(refs) or 'не определено'))
        for event in task['history']:
            outcome = 'согласовано' if event['accepted_in_dialogue'] else 'не согласовано'
            doc.add_paragraph(f"{event['kind']}: {event['previous_value'] or '—'} → {event['new_value'] or '—'} ({outcome})")
        refs = list(dict.fromkeys(ref for ids in task['evidence'].values() for ref in ids))
        for ref in refs:
            u = by_id[ref]
            doc.add_paragraph(f"[{ref}; {u['start_ms']/1000:g}–{u['end_ms']/1000:g} с] {u['text']}")
    footer = section.footer.paragraphs[0]
    footer.add_run('Локальный стенд · ' + report['source_sha256'][:16]).font.size = Pt(8)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
