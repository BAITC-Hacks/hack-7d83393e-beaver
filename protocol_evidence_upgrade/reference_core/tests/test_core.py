from __future__ import annotations
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import pytest
from docx import Document
from fastapi.testclient import TestClient
from evidence_core.audit import audit
from evidence_core.ledger import replay
from evidence_core.app import app
from evidence_core.export_docx import render

ROOT = Path(__file__).resolve().parents[1]
def load(name):
    return json.loads((ROOT/'fixtures'/name).read_text(encoding='utf-8'))
@pytest.fixture
def protocol(): return load('01_mixed_deadline.json')
@pytest.fixture
def ledger(): return load('ledger_mixed.json')
@pytest.fixture
def client():
    with TestClient(app) as c: yield c

@pytest.mark.parametrize('name',['01_mixed_deadline.json','02_ru_missing_fields.json','03_kk_no_new_task.json'])
def test_existing_contract_examples(name):
    assert audit(load(name))['status']=='PASS'

@pytest.mark.parametrize('mutation,code',[
    (lambda p:p['participants'].append(deepcopy(p['participants'][0])), 'DUPLICATE_ID'),
    (lambda p:p['tasks'][0].update(assignee_id='ghost'), 'UNKNOWN_ASSIGNEE'),
    (lambda p:p['tasks'][0]['evidence'].update(action=['missing']), 'BROKEN_EVIDENCE'),
    (lambda p:p['tasks'][0]['evidence'].update(assignee=[]), 'MISSING_FIELD_EVIDENCE'),
    (lambda p:p['tasks'][0]['evidence'].update(deadline=[]), 'MISSING_FIELD_EVIDENCE'),
    (lambda p:p['tasks'][0].update(due_date=None), 'TIME_WITHOUT_DATE'),
    (lambda p:p['utterances'][0].update(end_ms=0), 'TIME_RANGE'),
    (lambda p:p['speaker_map'][0].update(participant_id=None), 'EMPTY_CONFIRMED_MAPPING'),
    (lambda p:p['speaker_map'][0].update(participant_id='ghost'), 'UNKNOWN_PARTICIPANT'),
    (lambda p:p['utterances'][0].update(speaker_id='unknown'), 'UNMAPPED_VOICE'),
    (lambda p:p['tasks'][0].update(due_date='2026-02-30'), 'SCHEMA'),
    (lambda p:p['meeting'].update(timezone='Bad/Place'), 'TIMEZONE'),
    (lambda p:p['meeting'].update(started_at='2026-09-23T10:00:00'), 'SCHEMA'),
    (lambda p:p['summary'][0].update(evidence_ids=['missing']), 'BROKEN_EVIDENCE'),
    (lambda p:p['review_questions'].append({'task_id':'ghost','field':'action','text':'Уточнить'}), 'UNKNOWN_TASK'),
    (lambda p:p['tasks'][0]['history'].append(deepcopy(p['tasks'][0]['history'][0])), 'DUPLICATE_EVENT'),
    (lambda p:p.update(extra='not allowed'), 'SCHEMA'),
])
def test_auditor_rejects_structure_errors(protocol,mutation,code):
    mutation(protocol)
    report=audit(protocol)
    assert report['status']=='FAIL'
    assert code in {i['code'] for i in report['issues']}

@pytest.mark.parametrize('raw',[None,[],{},'string',42,True])
def test_auditor_handles_invalid_roots(raw): assert audit(raw)['status']=='FAIL'

def test_structure_never_claims_semantic_proof(protocol):
    protocol['tasks'][0]['action']='Совершенно неподтверждённое содержание'
    r=audit(protocol)
    assert r['status']=='PASS'  # References still exist; semantics deliberately not asserted.
    assert r['semantic_support']=='NOT RUN'
    assert r['asr']=='NOT RUN'

def test_hash_changes_with_input(protocol):
    original=audit(protocol)['source_sha256']
    protocol['meeting']['title']+=' changed'
    assert audit(protocol)['source_sha256']!=original

def test_audit_does_not_modify_input(protocol):
    before=deepcopy(protocol);audit(protocol);assert protocol==before

@pytest.mark.parametrize('at,expected_date,task_count',[
    (0,None,0),(5999,None,0),(6000,'2026-09-25',1),
    (16000,'2026-09-25',1),(21999,'2026-09-25',1),
    (22000,'2026-09-28',1),(35000,'2026-09-28',1),
    (40000,'2026-09-28',1),(48000,'2026-09-28',2),
])
def test_time_travel_is_prefix_correct(ledger,at,expected_date,task_count):
    result=replay(ledger,at)
    assert len(result['tasks'])==task_count
    assert all(e['at_ms']<=at for e in result['events'])
    if task_count:assert result['tasks'][0]['due_date']==expected_date

def test_agreement_not_proposal_changes_value(ledger):
    r=replay(ledger,16000)
    assert r['tasks'][0]['due_time'] is None
    assert len(r['tasks'][0]['pending_proposals'])==1
    assert replay(ledger,22000)['tasks'][0]['due_time']=='12:00:00'

def test_rejected_change_preserves_accepted_value(ledger):
    r=replay(ledger)
    assert r['tasks'][0]['due_date']=='2026-09-28'
    rejection=[e for e in r['events'] if e['kind']=='reject'][0]
    assert not rejection['changed_current_state']
    assert not r['tasks'][0]['pending_proposals']

def test_acceptance_combines_proposal_and_confirmation_evidence(ledger):
    task=replay(ledger)['tasks'][0]
    assert task['evidence']['due_date']==['u2','u3']
    assert task['evidence']['assignee_id']==['u1']

def test_unknowns_are_questions_not_invented_values(ledger):
    task=replay(ledger)['tasks'][1]
    assert task['assignee_id'] is None and task['due_date'] is None
    assert task['questions']==['Кто отвечает?','К какому сроку?']

def test_replay_deterministic_and_nonmutating(ledger):
    before=deepcopy(ledger)
    assert replay(ledger)==replay(ledger)
    assert ledger==before

@pytest.mark.parametrize('mutation',[
    lambda b:b['events'][0].update(evidence_ids=['ghost']),
    lambda b:b['events'][0].update(at_ms=1),
    lambda b:b['events'][0].update(task_id='other'),
    lambda b:b['events'][0]['initial'].update(assignee_id='ghost'),
    lambda b:b['events'][0]['initial'].update(due_date='2026-02-30'),
    lambda b:b['events'][0]['initial'].update(due_time='12:00:00+05:00'),
    lambda b:b['events'][0]['initial'].update(due_time='12:00:00.500000'),
    lambda b:b['events'][1].update(patch={'made_up':'field'}),
    lambda b:b['events'][2].update(proposal_id='missing'),
    lambda b:b['events'][4].update(proposal_id='e2'),
    lambda b:b['events'][0].update(event_id='e2'),
    lambda b:b['events'][1].update(at_ms=5500),
    lambda b:b['events'][0].update(at_ms=True),
    lambda b:b['events'][0]['initial'].update(action=''),
    lambda b:b['events'][1].update(patch={}),
    lambda b:b['utterances'][0].update(end_ms=0),
])
def test_ledger_rejects_invalid_events(ledger,mutation):
    mutation(ledger)
    with pytest.raises(ValueError):replay(ledger)

@pytest.mark.parametrize('cutoff',[-1,True,1.2,'12'])
def test_bad_cutoff_rejected(ledger,cutoff):
    with pytest.raises(ValueError):replay(ledger,cutoff)

def test_future_invalid_event_is_not_hidden_by_cutoff(ledger):
    ledger['events'][-1]['evidence_ids']=['missing']
    with pytest.raises(ValueError): replay(ledger,0)

def test_stale_proposal_cannot_overwrite_newer_agreement(ledger):
    stale=deepcopy(ledger['events'][1]);stale['event_id']='stale';stale['patch']['due_date']='2026-10-01'
    ledger['events'].insert(2,stale)
    ledger['events'].append({'event_id':'late-accept','task_id':'t1','kind':'accept','at_ms':48000,'evidence_ids':['u7'],'proposal_id':'stale'})
    with pytest.raises(ValueError,match='устарело'):replay(ledger)

def test_cancel_stops_active_task(ledger):
    ledger['events'].append({'event_id':'cancel','task_id':'t1','kind':'cancel','at_ms':48000,'evidence_ids':['u7']})
    task=replay(ledger)['tasks'][0]
    assert task['status']=='cancelled' and task['questions']==[]

def test_do_not_reopen_cancelled_task_implicitly(ledger):
    ledger['events'].extend([
        {'event_id':'cancel','task_id':'t1','kind':'cancel','at_ms':48000,'evidence_ids':['u7']},
        {'event_id':'change','task_id':'t1','kind':'propose','at_ms':48000,'evidence_ids':['u7'],'patch':{'due_date':'2026-10-01'}}])
    with pytest.raises(ValueError): replay(ledger)

def test_participant_rename_changes_ids_without_changing_logic(ledger):
    ledger['participants']=['x1','x2','x3']
    ledger['events'][0]['initial']['assignee_id']='x2'
    assert replay(ledger)['tasks'][0]['assignee_id']=='x2'

def test_export_roundtrip_matches_saved_deadline(protocol):
    doc=Document(BytesIO(render(protocol)))
    text='\n'.join(p.text for p in doc.paragraphs)
    assert 'СИНТЕТИЧЕСКИЙ' in text and 'Черновик' in text
    assert 'жұмаға' in text
    assert doc.tables[0].rows[1].cells[1].text=='Данияр'
    assert doc.tables[0].rows[1].cells[2].text=='2026-09-28 12:00:00'

def test_export_rejects_invalid_evidence(protocol):
    protocol['tasks'][0]['evidence']['action']=['ghost']
    with pytest.raises(ValueError):render(protocol)

def test_http_health_does_not_claim_model_readiness(client):
    response=client.get('/api/health')
    assert response.status_code==200
    assert response.json()['models_loaded'] is False

def test_http_capabilities_disclose_missing_components(client):
    assert 'asr' in client.get('/api/capabilities').json()['not_implemented']

def test_http_audit_and_docx(client,protocol):
    assert client.post('/api/audit',json=protocol).json()['status']=='PASS'
    export=client.post('/api/export.docx',json=protocol)
    assert export.status_code==200
    assert export.content[:2]==b'PK'

def test_http_replay(client,ledger):
    result=client.post('/api/replay',json={'bundle':ledger,'cutoff_ms':16000})
    assert result.status_code==200
    assert result.json()['tasks'][0]['due_date']=='2026-09-25'

def test_http_malformed_ledger(client):
    assert client.post('/api/replay',json={'bundle':{}}).status_code==422

def test_unknown_fixture_no_path_access(client):
    assert client.get('/api/fixtures/secrets').status_code==404

def test_cross_origin_write_denied(client,protocol):
    assert client.post('/api/audit',json=protocol,headers={'Origin':'https://untrusted.example'}).status_code==403

def test_payload_limit(client):
    assert client.post('/api/audit',content=b'x'*2_000_001,headers={'Content-Type':'application/json'}).status_code==413

def test_ui_static_local_security(client):
    index=client.get('/')
    assert index.status_code==200
    assert 'Синтетический стенд' in index.text
    assert 'http://' not in index.text and 'https://' not in index.text
    assert "frame-ancestors 'none'" in index.headers['Content-Security-Policy']
    js=client.get('/static/app.js').text
    assert 'textContent' in js and 'innerHTML' not in js

def test_no_external_swagger_assets(client):
    assert client.get('/docs').status_code==404
    assert client.get('/openapi.json').status_code==200
